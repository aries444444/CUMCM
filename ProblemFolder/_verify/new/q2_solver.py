# -*- coding: utf-8 -*-
"""2026 CUMCM A 题 问题2 求解器
整个烘干过程模型（预热平衡 + 恒温干燥两阶段烘房条件），经验公式统一采用附录3。
物性随 C、T 变，热质双向耦合，交错 Picard + Crank--Nicolson。
问题2：0--10800 s（主离散 Δr=0.125 mm, Δt=0.25 s）-> 表3/表4 + result2.xlsx（每 1 s、每 0.1 cm）。
烘房条件：0--14400 s 用附件1 插值，之后恒温段取附件1 末值 (50.165 °C, 0.04986)。
迭代：T、C 无穷范数相对变化 < 1e-8 停止（上限 20 轮），判据依据见 common.py。
验证：V1 时间收敛、V2 空间收敛（0.5/0.25/0.125 mm 同 Δt）、V3 渐近。
"""
import os
import numpy as np
import openpyxl

from common import ROOT, kernel, load_annex, picard_step, setup_console

setup_console()

R, dr, N = 0.02, 1.25e-4, 161          # 半径 2 cm，主离散 0.125 mm，161 节点
r = np.arange(N) * dr
T0, C0 = 28.0, 2.55
h, hm = 25.0, 8e-7
TEND, CEND = 50.165, 0.04986          # 附件1 末值（恒温干燥段）

# 附录3 经验式（问题2、问题3 统一采用）
rho23 = lambda C: 650.0 + 128.0 * C
cp23 = lambda C: 1450.0 + 2736.0 * C / (C + 1.0)
k23 = lambda C: 0.21 + 0.38 * C / (C + 1.0)
D23 = lambda C, T: 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / T)   # T 用开尔文

d = load_annex('附件1.xlsx')
T_air = lambda t: np.interp(t, d[:, 0], d[:, 1]) if t <= 14400 else TEND
C_air = lambda t: np.interp(t, d[:, 0], d[:, 2]) if t <= 14400 else CEND

def step(T, C, t_prev, t_cur, dt):
    """一个时间步：固定半径 + 附录3 物性的交错 Picard 耦合步。"""
    crk = kernel('radial', dt, N, dr)
    return picard_step(crk, T, C, dt, t_prev, t_cur,
                       rho23, cp23, k23, D23, T_air, C_air, h, hm)

def run3h(drg, Ng, dtt):
    """前 3 h 求解，返回 (表3, 表4, 迭代统计)。"""
    global N, dr, r
    N, dr, r = Ng, drg, np.arange(Ng) * drg
    T, C = np.full(N, T0), np.full(N, C0)
    t2 = np.arange(1800, 10801, 1800)
    ri = [0, Ng // 4, Ng // 2, 3 * Ng // 4, Ng - 1]
    tab3, tab4, iters = [], [], []
    for n in range(1, int(10800 / dtt) + 1):
        T, C, it, rel = step(T, C, (n - 1) * dtt, n * dtt, dtt)
        iters.append(it)
        if n * dtt in t2:
            tab3.append(T[ri].copy()); tab4.append(C[ri].copy())
    return np.array(tab3), np.array(tab4), np.array(iters)

# ================= 问题2：0--10800 s，主离散 Δr=0.125 mm, Δt=0.25 s =================
tab3, tab4, iters = run3h(1.25e-4, 161, 0.25)
tab3r = np.round(tab3, 4); tab4r = np.round(tab4, 4)
print('表3 温度 (°C)，行=0.5..3.0 h，列=0/0.5/1/1.5/2 cm:')
print(tab3r)
print('表4 水分浓度 (kg/kg):')
print(tab4r)
print('Picard 统计: 平均 %.2f 轮，最大 %d 轮，步数 %d' % (iters.mean(), iters.max(), len(iters)))

# result2.xlsx：每 1 s、每 0.1 cm（0.125 mm 离散上每 8 节点一列）
wb = openpyxl.Workbook()
wsT = wb.active; wsT.title = '温度'
wsC = wb.create_sheet('水分浓度')
xcols = list(range(0, N, 8))
for ws in (wsT, wsC):
    ws.cell(1, 1, '时间\\到药材中心的距离')
    for j, i in enumerate(xcols):
        ws.cell(1, 2 + j, round(r[i] * 100, 1))
T, C = np.full(N, T0), np.full(N, C0)
for n in range(1, 43201):
    T, C, it, rel = step(T, C, (n - 1) * 0.25, n * 0.25, 0.25)
    if n % 4 == 0:
        m = n // 4
        wsT.cell(1 + m, 1, m); wsC.cell(1 + m, 1, m)
        for j, i in enumerate(xcols):
            wsT.cell(1 + m, 2 + j, round(float(T[i]), 4))
            wsC.cell(1 + m, 2 + j, round(float(C[i]), 4))
wb.save(os.path.join(ROOT, 'result2.xlsx'))
print('已写出 result2.xlsx')

# ================= 验证 =================
# V1 时间收敛：Δt=0.5 vs 0.25（主离散，表3/表4 最大差）
tab3b, tab4b, _ = run3h(1.25e-4, 161, 0.5)
print('验证V1 时间收敛: 表3 最大差 %.2e °C，表4 最大差 %.2e kg/kg'
      % (np.max(np.abs(tab3b - tab3)), np.max(np.abs(tab4b - tab4))))

# V2 空间收敛：0.5 mm 与 0.25 mm 两套离散（同 Δt=0.25 s、同位置采样）
tab3_05, tab4_05, _ = run3h(5e-4, 41, 0.25)
tab3_025, tab4_025, _ = run3h(2.5e-4, 81, 0.25)
print('验证V2 空间收敛: 表3 最大差 0.5->0.25mm: %.2e °C, 0.25->0.125mm: %.2e °C'
      % (np.max(np.abs(tab3_05 - tab3_025)), np.max(np.abs(tab3_025 - tab3))))
print('              表4 最大差 0.5->0.25mm: %.2e, 0.25->0.125mm: %.2e kg/kg'
      % (np.max(np.abs(tab4_05 - tab4_025)), np.max(np.abs(tab4_025 - tab4))))

# V3 渐近：常数边界 (50.165, 0.04986)，长时程
N, dr, r = 21, 1e-3, np.arange(21) * 1e-3
T, C = np.full(N, T0), np.full(N, C0)
for n in range(1, int(3e5 / 2.5) + 1):
    T, C, it, rel = step(T, C, (n - 1) * 2.5, n * 2.5, 2.5)
    if n * 2.5 in (100000, 200000, 300000):
        print('验证V3 渐近: t=%ds 中心 T=%.4f (目标 %.4f)，表面 C=%.5f (目标 %.5f)'
              % (n * 2.5, T[0], TEND, C[-1], CEND))
