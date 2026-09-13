# -*- coding: utf-8 -*-
"""2026 CUMCM A 题 问题2 求解器
整个烘干过程模型（预热平衡 + 恒温干燥两阶段烘房条件），经验公式统一采用附录3。
物性随 C、T 变，热质双向耦合，交错 Picard + Crank--Nicolson + Thomas。
问题2：0--10800 s（主离散 Δr=0.125 mm, Δt=0.25 s）-> 表3/表4 + result2.xlsx（每 1 s、每 0.1 cm）。
烘房条件：0--14400 s 用附件1 插值，之后恒温段取附件1 末值 (50.165 °C, 0.04986)。
迭代：T、C 无穷范数相对变化 < 1e-8 停止（上限 20 轮）。
  依据：时间步误差量级实测为 O(1e-5)（Δt 减半差值），迭代残差须再低 2--3 个量级，
  取 1e-8（相对）使迭代误差不成为精度瓶颈；系数在迭代中取最新可用层。
验证：V1 时间收敛、V2 空间收敛（0.5/0.25/0.125 mm 同 Δt）、V3 渐近。
"""
import sys, os
import numpy as np
import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

ROOT = os.path.dirname(os.path.abspath(__file__))

def _annex(name):
    """定位赛题附件数据:优先脚本同目录(支撑材料布局),其次仓库目录结构。"""
    for p in (os.path.join(ROOT, name),
              os.path.join(ROOT, 'OriginalMaterial', 'A题', '附件', name)):
        if os.path.exists(p):
            return p
    raise FileNotFoundError('未找到 %s:请将赛题附件文件放在脚本同目录下' % name)

ANNEX1 = _annex('附件1.xlsx')

R, dr, N = 0.02, 1.25e-4, 161          # 半径 2 cm，主离散 0.125 mm，161 节点
r = np.arange(N) * dr
T0, C0 = 28.0, 2.55
h, hm = 25.0, 8e-7
TEND, CEND = 50.165, 0.04986          # 附件1 末值（恒温干燥段）
PICARD_TOL, PICARD_MAXIT = 1e-8, 20

# 附录3 经验式（问题2、问题3 统一采用）
rho23 = lambda C: 650.0 + 128.0 * C
cp23 = lambda C: 1450.0 + 2736.0 * C / (C + 1.0)
k23 = lambda C: 0.21 + 0.38 * C / (C + 1.0)
D23 = lambda C, T: 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / T)   # T 用开尔文

d = np.array([row for row in
              openpyxl.load_workbook(ANNEX1, data_only=True)['Sheet1'].iter_rows(values_only=True)][1:],
             float)
T_air = lambda t: np.interp(t, d[:, 0], d[:, 1]) if t <= 14400 else TEND
C_air = lambda t: np.interp(t, d[:, 0], d[:, 2]) if t <= 14400 else CEND

def thomas(a, b, c, rhs):
    n = len(rhs); cp = np.zeros(n); dp = np.zeros(n)
    cp[0] = c[0] / b[0]; dp[0] = rhs[0] / b[0]
    for i in range(1, n):
        m = b[i] - a[i] * cp[i - 1]; cp[i] = c[i] / m; dp[i] = (rhs[i] - a[i] * dp[i - 1]) / m
    x = np.zeros(n); x[-1] = dp[-1]
    for i in range(n - 2, -1, -1): x[i] = dp[i] - cp[i] * x[i + 1]
    return x

def crank_step(u, K, Cp, beta, bc_prev, bc_next, dt):
    """Cp_i (u^{n+1}_i - u^n_i)/dt = Lu + s；K: 每节点扩散系数；Cp: 每节点容量系数。
    热方程: K=k(C), Cp=rho(C)c_p(C), beta=h/k_s；质方程: K=D, Cp=1, beta=h_m/D_s。"""
    Km = (K[:-1] + K[1:]) / 2
    lo = np.zeros(N); di = np.zeros(N); up = np.zeros(N); s = np.zeros(N)
    di[0] = -4 * Km[0] / dr ** 2; up[0] = 4 * Km[0] / dr ** 2
    for i in range(1, N - 1):
        lo[i] = Km[i - 1] * (i - 0.5) / (i * dr ** 2)
        di[i] = -(Km[i] * (i + 0.5) + Km[i - 1] * (i - 0.5)) / (i * dr ** 2)
        up[i] = Km[i] * (i + 0.5) / (i * dr ** 2)
    # 表面控制体 [r_{N-3/2}, R]：表面节点 r_{N-1}=(N-1)dr=R 位于真实表面，
    # Robin 条件直接取节点值：表面通量 -R·beta·K_s·(u_{N-1}-bc)
    # （热: -R·h(T_R-T_air)；质: -R·h_m(C_R-C_air)），与内部界面通量逐时间层严格守恒。
    vhat = ((N - 1) ** 2 - (N - 1.5) ** 2) * dr ** 2 / 2
    Bc = (N - 1) * dr * beta * K[-1]
    lo[N - 1] = Km[N - 2] * (N - 1.5) / vhat
    di[N - 1] = -(lo[N - 1] + Bc / vhat)
    up[N - 1] = 0.0
    def src(bc):
        s = np.zeros(N); s[N - 1] = Bc * bc / vhat; return s
    # Crank--Nicolson: Cp_i (u^{n+1}-u^n)/dt = 1/2(Lu^{n+1}+Lu^n) + 1/2(s^{n+1}+s^n)
    lhs_lo = -0.5 * lo
    lhs_di = Cp / dt - 0.5 * di
    lhs_up = -0.5 * up
    rhs = Cp / dt * u + 0.5 * (lo * np.roll(u, 1) + di * u + up * np.roll(u, -1)) \
        + 0.5 * (src(bc_next) + src(bc_prev))
    rhs[0] = Cp[0] / dt * u[0] + 0.5 * (di[0] * u[0] + up[0] * u[1]) \
        + 0.5 * (src(bc_next)[0] + src(bc_prev)[0])
    rhs[N - 1] = Cp[N - 1] / dt * u[N - 1] \
        + 0.5 * (lo[N - 1] * u[N - 2] + di[N - 1] * u[N - 1]) \
        + 0.5 * (src(bc_next)[N - 1] + src(bc_prev)[N - 1])
    return thomas(lhs_lo, lhs_di, lhs_up, rhs)

def step(T, C, t_prev, t_cur, dt):
    """一个时间步：交错 Picard，每轮从 u^n 出发；T、C 相对变化（无穷范数，分母取
    max(‖u‖∞,1)）小于 1e-8 或满 20 轮停止。返回 (T, C, 轮数, 退出残差)。"""
    Tn, Cn = T.copy(), C.copy()
    for it in range(1, PICARD_MAXIT + 1):
        Kk = k23(C); Rc = rho23(C) * cp23(C)
        T = crank_step(Tn, Kk, Rc, h / Kk[-1], T_air(t_prev), T_air(t_cur), dt)
        Tk = T + 273.15
        Dc = D23(C, Tk)
        C = crank_step(Cn, Dc, np.ones(N), hm / Dc[-1], C_air(t_prev), C_air(t_cur), dt)
        if it > 1:
            rel = max(np.max(np.abs(T - To)) / max(np.max(np.abs(T)), 1.0),
                      np.max(np.abs(C - Co)) / max(np.max(np.abs(C)), 1.0))
            if rel < PICARD_TOL:
                return T, C, it, rel
        To, Co = T.copy(), C.copy()
    return T, C, PICARD_MAXIT, rel

def run3h(drg, Ng, dtt):
    """前 3 h 求解，返回 (表3, 表4, 迭代统计)。"""
    global N, dr, r
    N, dr, r = Ng, drg, np.arange(Ng) * drg
    T, C = np.full(N, T0), np.full(N, C0)
    t2 = np.arange(1800, 10801, 1800)
    ri = [0, Ng // 4, Ng // 2, 3 * Ng // 4, Ng - 1]
    tab3, tab4, iters = [], [], []
    for n in range(1, int(10800 / dtt) + 1):
        T, C, it, rel = step(T, C, (n - 1) * dtt, n * dtt, dtt)
        iters.append(it)
        if n * dtt in t2:
            tab3.append(T[ri].copy()); tab4.append(C[ri].copy())
    return np.array(tab3), np.array(tab4), np.array(iters)

# ================= 问题2：0--10800 s，主离散 Δr=0.125 mm, Δt=0.25 s =================
tab3, tab4, iters = run3h(1.25e-4, 161, 0.25)
tab3r = np.round(tab3, 4); tab4r = np.round(tab4, 4)
print('表3 温度 (°C)，行=0.5..3.0 h，列=0/0.5/1/1.5/2 cm:')
print(tab3r)
print('表4 水分浓度 (kg/kg):')
print(tab4r)
print('Picard 统计: 平均 %.2f 轮，最大 %d 轮，步数 %d' % (iters.mean(), iters.max(), len(iters)))

# result2.xlsx：每 1 s、每 0.1 cm（0.125 mm 离散上每 8 节点一列）
wb = openpyxl.Workbook()
wsT = wb.active; wsT.title = '温度'
wsC = wb.create_sheet('水分浓度')
xcols = list(range(0, N, 8))
for ws in (wsT, wsC):
    ws.cell(1, 1, '时间\\到药材中心的距离')
    for j, i in enumerate(xcols):
        ws.cell(1, 2 + j, round(r[i] * 100, 1))
T, C = np.full(N, T0), np.full(N, C0)
for n in range(1, 43201):
    T, C, it, rel = step(T, C, (n - 1) * 0.25, n * 0.25, 0.25)
    if n % 4 == 0:
        m = n // 4
        wsT.cell(1 + m, 1, m); wsC.cell(1 + m, 1, m)
        for j, i in enumerate(xcols):
            wsT.cell(1 + m, 2 + j, round(float(T[i]), 4))
            wsC.cell(1 + m, 2 + j, round(float(C[i]), 4))
wb.save(os.path.join(ROOT, 'result2.xlsx'))
print('已写出 result2.xlsx')

# ================= 验证 =================
# V1 时间收敛：Δt=0.5 vs 0.25（主离散，表3/表4 最大差）
tab3b, tab4b, _ = run3h(1.25e-4, 161, 0.5)
print('验证V1 时间收敛: 表3 最大差 %.2e °C，表4 最大差 %.2e kg/kg'
      % (np.max(np.abs(tab3b - tab3)), np.max(np.abs(tab4b - tab4))))

# V2 空间收敛：0.5 mm 与 0.25 mm 两套离散（同 Δt=0.25 s、同位置采样）
tab3_05, tab4_05, _ = run3h(5e-4, 41, 0.25)
tab3_025, tab4_025, _ = run3h(2.5e-4, 81, 0.25)
print('验证V2 空间收敛: 表3 最大差 0.5->0.25mm: %.2e °C, 0.25->0.125mm: %.2e °C'
      % (np.max(np.abs(tab3_05 - tab3_025)), np.max(np.abs(tab3_025 - tab3))))
print('              表4 最大差 0.5->0.25mm: %.2e, 0.25->0.125mm: %.2e kg/kg'
      % (np.max(np.abs(tab4_05 - tab4_025)), np.max(np.abs(tab4_025 - tab4))))

# V3 渐近：常数边界 (50.165, 0.04986)，长时程
N, dr, r = 21, 1e-3, np.arange(21) * 1e-3
T, C = np.full(N, T0), np.full(N, C0)
for n in range(1, int(3e5 / 2.5) + 1):
    T, C, it, rel = step(T, C, (n - 1) * 2.5, n * 2.5, 2.5)
    if n * 2.5 in (100000, 200000, 300000):
        print('验证V3 渐近: t=%ds 中心 T=%.4f (目标 %.4f)，表面 C=%.5f (目标 %.5f)'
              % (n * 2.5, T[0], TEND, C[-1], CEND))
