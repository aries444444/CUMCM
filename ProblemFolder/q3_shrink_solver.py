# -*- coding: utf-8 -*-
"""2026 CUMCM A 题 消融对照求解器（审稿意见 6 的 2×2 表）

同一离散设置（ζ 网格 Δζ=1/160、Δt=5 s、Picard 相对残差 < 1e-8）下，
物性（附录3 / 附录4）x 几何（固定半径 R=2 cm / 附件2 收缩 R(t)）四个组合的烘干时长：

             固定半径      收缩半径
  附录3      t3,F          t3,S
  附录4      t4,F          t4,S

依据：问题 3 与问题 4 同时更换了物性与几何，只有同物性比较固定/收缩半径，
才能分离收缩的净效应（审稿意见 6）。固定半径情形 ζ 系方程即径向方程（R≡R0）。
"""
import sys, os
import numpy as np
import openpyxl
from scipy.linalg import solve_banded
from scipy.interpolate import PchipInterpolator

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

ROOT = os.path.dirname(os.path.abspath(__file__))
ANNEX1 = os.path.join(ROOT, 'OriginalMaterial', 'A题', '附件', '附件1.xlsx')
ANNEX2 = os.path.join(ROOT, 'OriginalMaterial', 'A题', '附件', '附件2.xlsx')

R0, T0, C0 = 0.02, 28.0, 2.55
h, hm = 25.0, 8e-7
TEND, CEND = 50.165, 0.04986
CTARGET = 0.15
DZ, NZ = 1.0 / 160.0, 161
zeta = np.arange(NZ) * DZ
PICARD_TOL, PICARD_MAXIT = 1e-8, 20
DT = 5.0

rho3 = lambda C: 650.0 + 128.0 * C
cp3 = lambda C: 1450.0 + 2736.0 * C / (C + 1.0)
k3 = lambda C: 0.21 + 0.38 * C / (C + 1.0)
D3 = lambda C, T: 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / T)
rho4 = lambda C: 760.0 + 90.0 * C
cp4 = lambda C: 1850.0 + 2150.0 * C / (C + 1.0)
k4 = lambda C: 0.12 + 0.20 * C / (C + 1.0)
D4 = lambda C, T: 4.2e-4 * np.exp(-0.30 / np.maximum(C, 0.02)) * np.exp(-3850.0 / T)

d1 = np.array([row for row in
               openpyxl.load_workbook(ANNEX1, data_only=True)['Sheet1'].iter_rows(values_only=True)][1:],
              float)
T_air = lambda t: np.interp(t, d1[:, 0], d1[:, 1]) if t <= 14400 else TEND
C_air = lambda t: np.interp(t, d1[:, 0], d1[:, 2]) if t <= 14400 else CEND

d2 = np.array([row for row in
               openpyxl.load_workbook(ANNEX2, data_only=True)['Sheet1'].iter_rows(values_only=True)][1:],
              float)
T2END = d2[-1, 0]; REND = d2[-1, 1] / 100.0
_Rp = PchipInterpolator(d2[:, 0], d2[:, 1] / 100.0, extrapolate=False)
_Rpd = _Rp.derivative()

def thomas(a, b, c, rhs):
    n = len(rhs)
    ab = np.zeros((3, n))
    ab[0, 1:] = c[:-1]
    ab[1, :] = b
    ab[2, :-1] = a[1:]
    return solve_banded((1, 1), ab, rhs)

def crank_step(u, K, Cp, R, beta, bc_prev, bc_next, dt):
    Km = (K[:-1] + K[1:]) / 2
    R2 = R * R
    lo = np.zeros(NZ); di = np.zeros(NZ); up = np.zeros(NZ)
    di[0] = -4 * Km[0] / (DZ ** 2 * R2); up[0] = 4 * Km[0] / (DZ ** 2 * R2)
    for i in range(1, NZ - 1):
        lo[i] = Km[i - 1] * (i - 0.5) / (i * DZ ** 2 * R2)
        di[i] = -(Km[i] * (i + 0.5) + Km[i - 1] * (i - 0.5)) / (i * DZ ** 2 * R2)
        up[i] = Km[i] * (i + 0.5) / (i * DZ ** 2 * R2)
    vhat = (1.0 - ((NZ - 1.5) * DZ) ** 2) / 2.0
    Bc = beta * K[-1]
    lo[NZ - 1] = Km[NZ - 2] * (NZ - 1.5) / (vhat * R2)
    di[NZ - 1] = -(lo[NZ - 1] + Bc / (vhat * R2))
    up[NZ - 1] = 0.0
    def src(bc):
        s = np.zeros(NZ); s[NZ - 1] = Bc * bc / (vhat * R2); return s
    lhs_lo = -0.5 * lo
    lhs_di = Cp / dt - 0.5 * di
    lhs_up = -0.5 * up
    rhs = Cp / dt * u + 0.5 * (lo * np.roll(u, 1) + di * u + up * np.roll(u, -1)) \
        + 0.5 * (src(bc_next) + src(bc_prev))
    rhs[0] = Cp[0] / dt * u[0] + 0.5 * (di[0] * u[0] + up[0] * u[1]) \
        + 0.5 * (src(bc_next)[0] + src(bc_prev)[0])
    rhs[NZ - 1] = Cp[NZ - 1] / dt * u[NZ - 1] \
        + 0.5 * (lo[NZ - 1] * u[NZ - 2] + di[NZ - 1] * u[NZ - 1]) \
        + 0.5 * (src(bc_next)[NZ - 1] + src(bc_prev)[NZ - 1])
    return thomas(lhs_lo, lhs_di, lhs_up, rhs)

def run_case(props, shrink):
    """props: 'A3' 或 'A4'；shrink: True 用附件2 R(t)，False 固定 R=R0。返回 (t*, 统计)。"""
    P = (rho3, cp3, k3, D3) if props == 'A3' else (rho4, cp4, k4, D4)
    rhoF, cpF, kF, DF = P
    if shrink:
        RofT = lambda t: float(_Rp(t)) if t <= T2END else REND
    else:
        RofT = lambda t: R0
    T, C = np.full(NZ, T0), np.full(NZ, C0)
    iters, rels = [], []
    nmax = int(300 * 3600 / DT)
    tstar = None
    for n in range(1, nmax + 1):
        tm = (n - 0.5) * DT
        R = RofT(tm)
        Cprev = C.copy()
        Tn, Cn = T.copy(), C.copy()
        rel = 0.0
        for it in range(1, PICARD_MAXIT + 1):
            Kk = kF(C); Rc = rhoF(C) * cpF(C)
            T = crank_step(Tn, Kk, Rc, R, h * R / Kk[-1], T_air((n - 1) * DT), T_air(n * DT), DT)
            Tk = T + 273.15
            Dc = DF(C, Tk)
            C = crank_step(Cn, Dc, np.ones(NZ), R, hm * R / Dc[-1], C_air((n - 1) * DT), C_air(n * DT), DT)
            if it > 1:
                rel = max(np.max(np.abs(T - To)) / max(np.max(np.abs(T)), 1.0),
                          np.max(np.abs(C - Co)) / max(np.max(np.abs(C)), 1.0))
                if rel < PICARD_TOL:
                    break
            To, Co = T.copy(), C.copy()
        iters.append(it); rels.append(rel)
        if C.max() <= CTARGET:
            gp = Cprev.max() - CTARGET
            gn = C.max() - CTARGET
            frac = gp / (gp - gn)
            tstar = (n - 1) * DT + DT * frac
            break
    if tstar is None:
        raise RuntimeError('300 h 内未达烘干判据: %s shrink=%s' % (props, shrink))
    return tstar, np.array(iters), np.array(rels)

if __name__ == "__main__":
    print('2×2 消融对照（Δζ=1/160, Δt=5 s, Picard<1e-8）:')
    t3F, i3F, r3F = run_case('A3', False)
    t3S, i3S, r3S = run_case('A3', True)
    t4F, i4F, r4F = run_case('A4', False)
    t4S, i4S, r4S = run_case('A4', True)
    for name, t, it, rl in (('t3,F', t3F, i3F, r3F), ('t3,S', t3S, i3S, r3S),
                            ('t4,F', t4F, i4F, r4F), ('t4,S', t4S, i4S, r4S)):
        print('%s = %.4f h (%d s), 平均 %.2f 轮 / 最大 %d 轮'
              % (name, t / 3600, round(t), it.mean(), it.max()))
    print('收缩净效应: 附录3 下 Δt* = %.4f h（%.1f%%），附录4 下 Δt* = %.4f h（%.1f%%）'
          % ((t3S - t3F) / 3600, (t3S / t3F - 1) * 100, (t4S - t4F) / 3600, (t4S / t4F - 1) * 100))
