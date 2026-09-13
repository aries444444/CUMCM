# -*- coding: utf-8 -*-
"""
2026 CUMCM A题 · 药材的烘干问题 —— 数据可视化
================================================
数据源（自动从上级 ProblemFolder 读取，无需手工粘贴）:
  ProblemFolder/result3.xlsx                            问题3 水分浓度 (每60 s, 0~2 cm)
  ProblemFolder/result4.xlsx                            问题4 水分浓度 (每60 s, 0~2 cm + 药材表面)
  ProblemFolder/OriginalMaterial/A题/附件/附件2.xlsx     收缩半径 R(t)

输出（本目录）:
  fig1_overview.png     总览四联图（中心对比 / 收缩 / 位置演化 / 径向剖面）
  fig2_heatmap.png      含水率时空热力图（问题3 vs 问题4，含动边界）
  fig3_validation.png   模型检验汇总（误差量级 + 烘干时长对比）

运行:  python visualize.py
"""
import os
import numpy as np
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter, NullFormatter

# ---------------- 全局样式 ----------------
plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei']
plt.rcParams['axes.unicode_minus'] = False
plt.rcParams['mathtext.fontset'] = 'dejavusans'   # 避免 SimHei 缺数学字形告警
plt.rcParams['axes.grid'] = True
plt.rcParams['grid.alpha'] = 0.3
plt.rcParams['grid.linestyle'] = '--'
plt.rcParams['axes.axisbelow'] = True
plt.rcParams['figure.dpi'] = 120
plt.rcParams['savefig.dpi'] = 300
plt.rcParams['savefig.bbox'] = 'tight'

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(HERE)
PF = os.path.join(PROJ, 'ProblemFolder')
ANNEX2 = os.path.join(PF, 'OriginalMaterial', 'A题', '附件', '附件2.xlsx')

C0 = 2.55          # 初始干基含水率 kg/kg
CTARGET = 0.15     # 烘干判据 kg/kg
TSTAR3, TSTAR4 = 205701.0, 182897.5   # 烘干时长 (s)，生产值（论文 6.3/6.4）

# 统一配色
CLR3, CLR4 = '#2563eb', '#dc2626'     # 问题3 蓝 / 问题4 红
CLR_SURF = '#059669'                  # 表面 绿
CLR_R = '#7c3aed'                     # 半径 紫
CLR_TARGET = '#f59e0b'                # 判据线 橙


# ---------------- 数据读取 ----------------
def read_result(path):
    """读取 result*.xlsx -> (t[s], dist[cm], C[nt,nd], C_surface[nt] or None)"""
    rows = list(openpyxl.load_workbook(path, data_only=True, read_only=True)
                .active.iter_rows(values_only=True))
    header = list(rows[0])
    ncol = len(header)
    dist_cols = [j for j in range(1, ncol) if isinstance(header[j], (int, float))]
    surf_col = ncol - 1 if isinstance(header[ncol - 1], str) else None
    dist = np.array([header[j] for j in dist_cols], float)
    t = np.array([r[0] for r in rows[1:]], float)
    C = np.array([[np.nan if r[j] is None else float(r[j]) for j in dist_cols]
                  for r in rows[1:]], float)
    Cs = (np.array([np.nan if r[surf_col] is None else float(r[surf_col])
                    for r in rows[1:]], float) if surf_col is not None else None)
    return t, dist, C, Cs


def read_radius():
    """读取附件2 -> (t[s], R[cm])"""
    rows = list(openpyxl.load_workbook(ANNEX2, data_only=True, read_only=True)
                .active.iter_rows(values_only=True))[1:]
    d = np.array(rows, float)
    return d[:, 0], d[:, 1]          # s, cm


t3, d3, C3, _ = read_result(os.path.join(PF, 'result3.xlsx'))
t4, d4, C4, Cs4 = read_result(os.path.join(PF, 'result4.xlsx'))
tR, Rcm = read_radius()

h3, h4 = t3 / 3600.0, t4 / 3600.0
hR = tR / 3600.0
R_at = lambda tt: np.interp(tt, tR, Rcm)          # tt in s -> cm
i0, i05, i10 = 0, int(np.argmin(abs(d3 - 0.5))), int(np.argmin(abs(d3 - 1.0)))

print('问题3: %d 行, %d 列, 0~%.2f h' % (len(t3), len(d3), h3[-1]))
print('问题4: %d 行, %d 列 + 表面, 0~%.2f h' % (len(t4), len(d4), h4[-1]))
print('附件2 半径: %.3f -> %.3f cm' % (Rcm[0], Rcm[-1]))


# ================= 图 1: 总览四联图 =================
fig, axes = plt.subplots(2, 2, figsize=(14.5, 10))
fig.suptitle('A题 药材烘干 · 含收缩前后对比总览', fontsize=17, fontweight='bold', y=0.985)

# -- (a) 中心含水率对比 + 判据线 + 交叉点 --
ax = axes[0, 0]
ax.plot(h3, C3[:, i0], color=CLR3, lw=2, label='问题3（无收缩, 附录3）')
ax.plot(h4, C4[:, i0], color=CLR4, lw=2, label='问题4（含收缩, 附录4）')
ax.axhline(CTARGET, color=CLR_TARGET, lw=1.8, ls='--', label='判据 $C=0.15$ kg/kg')
ax.axvline(h3[-1], color=CLR3, lw=1, ls=':', alpha=0.8)
ax.axvline(h4[-1], color=CLR4, lw=1, ls=':', alpha=0.8)
ax.plot([h3[-1]], [C3[-1, i0]], 'o', color=CLR3, ms=8, zorder=5)
ax.plot([h4[-1]], [C4[-1, i0]], 'o', color=CLR4, ms=8, zorder=5)
ax.annotate('$t^*=57.14$ h', (h3[-1], C3[-1, i0]), xytext=(-112, 20),
            textcoords='offset points', color=CLR3, fontsize=10,
            arrowprops=dict(arrowstyle='->', color=CLR3, lw=1.2))
ax.annotate('$t^*=50.80$ h', (h4[-1], C4[-1, i0]), xytext=(14, 74),
            textcoords='offset points', color=CLR4, fontsize=10,
            arrowprops=dict(arrowstyle='->', color=CLR4, lw=1.2))
# 交叉点（跳过起始段的伪交叉）
tg = np.linspace(0, h4[-1], 4000)
diff = np.interp(tg, h3, C3[:, i0]) - np.interp(tg, h4, C4[:, i0])
k = [i for i in np.where(np.diff(np.sign(diff)) != 0)[0] if tg[i] > 5]
if k:
    hx = tg[k[-1]]
    hy = float(np.interp(hx, h4, C4[:, i0]))
    ax.plot([hx], [hy], 'o', color='#334155', ms=7, zorder=6)
    ax.annotate('约 %.0f h 交叉\n（前期问题4更慢，后期反超）' % hx, (hx, hy),
                xytext=(-25, 92), textcoords='offset points', fontsize=9,
                color='#334155', ha='center',
                arrowprops=dict(arrowstyle='->', color='#334155', lw=1,
                                connectionstyle='arc3,rad=0.25'))
ax.set_xlabel('时间 / h'); ax.set_ylabel('中心含水率 / (kg/kg)')
ax.set_title('(a) 中心含水率随时间演化', fontsize=12, fontweight='bold')
ax.legend(fontsize=9, loc='upper right'); ax.set_xlim(0, 60); ax.set_ylim(0, 2.7)

# -- (b) 收缩半径 R(t) + 表面含水率 --
ax = axes[0, 1]
ax.plot(hR, Rcm, color=CLR_R, lw=2.5, label='半径 $R(t)$（附件2）')
ax.set_xlabel('时间 / h'); ax.set_ylabel('药材半径 / cm', color=CLR_R)
ax.tick_params(axis='y', labelcolor=CLR_R); ax.set_ylim(1.1, 2.1)
ax.axvline(67, color=CLR_R, lw=1, ls=':', alpha=0.8)
ax.annotate('67 h 后稳定于 1.198 cm', (67, 1.9), xytext=(78, 1.93),
            fontsize=9, color=CLR_R,
            arrowprops=dict(arrowstyle='->', color=CLR_R, lw=1))
ax2 = ax.twinx()
ax2.plot(h4, Cs4, color=CLR_SURF, lw=2, label='表面含水率（问题4）')
ax2.set_ylabel('表面含水率 / (kg/kg)', color=CLR_SURF)
ax2.tick_params(axis='y', labelcolor=CLR_SURF); ax2.set_ylim(0, 2.7); ax2.grid(False)
ln1, lb1 = ax.get_legend_handles_labels(); ln2, lb2 = ax2.get_legend_handles_labels()
ax.legend(ln1 + ln2, lb1 + lb2, fontsize=9, loc='center right')
ax.set_title('(b) 尺寸收缩与表面失水（收缩 40.1%）', fontsize=12, fontweight='bold')

# -- (c) 问题4 各位置含水率演化 --
ax = axes[1, 0]
for idx, lab, c in [(i0, '中心 $d=0$', '#b91c1c'), (i05, '$d=0.5$ cm', '#ea580c'),
                    (i10, '$d=1.0$ cm', '#ca8a04')]:
    ax.plot(h4, C4[:, idx], lw=2, color=c, label=lab)
ax.plot(h4, Cs4, lw=2, color=CLR_SURF, label='药材表面')
ax.axhline(CTARGET, color=CLR_TARGET, lw=1.8, ls='--')
ax.text(2, 0.17, '判据 0.15', color=CLR_TARGET, fontsize=9)
# 1.5 cm 何时移出药材（2.0 cm 自初始即为表面）
t15 = float(np.interp(1.5, Rcm[::-1], tR[::-1])) / 3600.0
ax.axvline(t15, color='#94a3b8', lw=1.2, ls='-.')
ax.text(4.5, 2.32, '灰色点划线：$d=1.5$ cm 处 %.1f h 后移出药材\n（$d=2.0$ cm 自初始即为表面）' % t15,
        fontsize=9, color='#475569', va='top')
ax.set_xlabel('时间 / h'); ax.set_ylabel('含水率 / (kg/kg)')
ax.set_title('(c) 问题4 各位置含水率演化', fontsize=12, fontweight='bold')
ax.legend(fontsize=9); ax.set_xlim(0, 52)

# -- (d) 问题4 径向剖面（含动边界） --
ax = axes[1, 1]
dgrid = np.linspace(0, 2.0, 400)
times = [6, 12, 24, 36, 50.80]
cmap = plt.get_cmap('plasma')
for j, th in enumerate(times):
    i = int(np.argmin(abs(h4 - th)))
    row = C4[i]
    valid = ~np.isnan(row)
    xv, yv = d4[valid], row[valid]
    Rv = R_at(t4[i])
    xx = dgrid[dgrid <= Rv]
    ax.plot(xx, np.interp(xx, xv, yv), lw=2, color=cmap(j / max(len(times) - 1, 1)),
            label='%.2f h' % h4[i])
    if j in (0, 3, 4):
        ax.axvline(Rv, color=cmap(j / max(len(times) - 1, 1)), lw=1, ls=':')
ax.axvline(2.0, color='#94a3b8', lw=1.5, ls='--')
ax.text(2.02, 2.3, '初始边界 2.0 cm', fontsize=8, color='#64748b', rotation=90, va='top')
ax.annotate('虚线 = 该时刻实际表面 $R(t)$', (1.4, 1.55), fontsize=9, color='#334155')
ax.set_xlabel('到药材中心的距离 / cm'); ax.set_ylabel('含水率 / (kg/kg)')
ax.set_title('(d) 问题4 径向剖面随时间的收缩', fontsize=12, fontweight='bold')
ax.legend(fontsize=9, title='时刻', ncol=2); ax.set_xlim(0, 2.35)

fig.tight_layout(rect=[0, 0, 1, 0.97])
p1 = os.path.join(HERE, 'fig1_overview.png')
fig.savefig(p1)
fig.savefig(os.path.join(PF, 'fig1_overview_solution.png'))
plt.close(fig)
print('已保存', p1)


# ================= 图 2: 时空热力图 =================
dgrid = np.linspace(0, 2.0, 201)


def build_Z(h, dist, C, shrink=False):
    Z = np.full((len(h), len(dgrid)), np.nan)
    for i in range(len(h)):
        row = C[i]
        valid = ~np.isnan(row)
        xv, yv = dist[valid], row[valid]
        if len(xv) < 2:
            continue
        lim = xv[-1] if shrink else 2.0
        m = dgrid <= lim
        Z[i, m] = np.interp(dgrid[m], xv, yv)
    return Z


Z3 = build_Z(h3, d3, C3, shrink=False)
Z4 = build_Z(h4, d4, C4, shrink=True)

fig, axes = plt.subplots(1, 2, figsize=(15, 5.6), sharey=True)
fig.suptitle('含水率时空分布 $C(d,t)$：收缩使有效扩散域不断缩小', fontsize=15,
             fontweight='bold', y=0.99)
vmin, vmax = 0.0, 2.55
for ax, (Z, h, title, tstar) in zip(
        axes, [(Z3, h3, '(a) 问题3 · 无收缩（$R\\equiv2$ cm）', TSTAR3 / 3600),
               (Z4, h4, '(b) 问题4 · 含收缩（$R(t)$）', TSTAR4 / 3600)]):
    pm = ax.pcolormesh(h, dgrid, Z.T, cmap='YlGnBu', vmin=vmin, vmax=vmax,
                       shading='auto', rasterized=True)
    ax.set_xlabel('时间 / h'); ax.set_title(title, fontsize=12, fontweight='bold')
    ax.set_xlim(0, 60)          # 统一横轴范围，使 t* 竖线不贴边界被裁掉
    ax.grid(False)
axes[0].set_ylabel('到药材中心的距离 / cm')
axes[1].plot(h4, [R_at(x * 3600) for x in h4], color='#dc2626', lw=2.2,
             label='实际表面 $R(t)$')
axes[1].axvline(TSTAR4 / 3600, color='#dc2626', lw=2.2, ls='--', label='$t^*=50.80$ h')
axes[0].axvline(TSTAR3 / 3600, color='#2563eb', lw=2.2, ls='--', label='$t^*=57.14$ h')
axes[0].legend(fontsize=9, loc='upper right'); axes[1].legend(fontsize=9, loc='upper right')
axes[1].axhline(1.198, color='#dc2626', lw=1, ls=':', alpha=0.7)
cb = fig.colorbar(pm, ax=axes, orientation='vertical', pad=0.02, aspect=28)
cb.set_label('含水率 / (kg/kg)')
p2 = os.path.join(HERE, 'fig2_heatmap.png')
fig.savefig(p2); plt.close(fig)
print('已保存', p2)


# ================= 图 3: 模型检验汇总 =================
fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.2),
                         gridspec_kw={'width_ratios': [1.45, 1]})
fig.suptitle('模型检验：误差量级与烘干时长', fontsize=15, fontweight='bold', y=0.99)

# 数据来源: ProblemFolder/q1~q4_out.txt
labels = ['问题1\n守恒闭合差', '问题3\n守恒闭合差', '问题4\n守恒闭合差',
          '问题3\nV2 时间收敛', '问题4\nV2 时间收敛',
          '问题3\nV3 空间收敛', '问题4\nV3 空间收敛']
vals = [1.5e-18, 1.7e-17, 4.15e-7, 5.58e-6, 1.10e-5, 5.27e-4, 5.33e-4]
cols = [CLR3, CLR3, CLR4, CLR3, CLR4, CLR3, CLR4]
ax = axes[0]
bars = ax.bar(range(len(vals)), vals, color=cols, alpha=0.85, edgecolor='white')
ax.set_yscale('log'); ax.set_ylim(1e-19, 1e-3)
# 纯文本刻度，避免 SimHei 缺 U+2212（减号）字形
ax.yaxis.set_major_formatter(FuncFormatter(
    lambda v, p: '' if v <= 0 else '1e%d' % int(round(np.log10(v)))))
ax.yaxis.set_minor_formatter(NullFormatter())
ax.axhline(1e-4, color=CLR_TARGET, lw=2, ls='--')
ax.text(len(vals) - 0.5, 1.3e-4, '四位小数目标 $10^{-4}$', color=CLR_TARGET,
        fontsize=9.5, ha='right')
for b, v in zip(bars, vals):
    ax.text(b.get_x() + b.get_width() / 2, v * 1.6, '%.2e' % v,
            ha='center', fontsize=8, rotation=0)
ax.set_xticks(range(len(vals))); ax.set_xticklabels(labels, fontsize=8.5)
ax.set_ylabel('误差量级（对数轴）')
ax.set_title('(a) 各项检验的误差量级', fontsize=12, fontweight='bold')

ax = axes[1]
bars = ax.bar(['问题3\n（无收缩）', '问题4\n（含收缩）'], [TSTAR3 / 3600, TSTAR4 / 3600],
              color=[CLR3, CLR4], alpha=0.85, width=0.55, edgecolor='white')
for b, v in zip(bars, [TSTAR3 / 3600, TSTAR4 / 3600]):
    ax.text(b.get_x() + b.get_width() / 2, v + 0.8, '%.2f h' % v,
            ha='center', fontsize=11, fontweight='bold')
ax.annotate('', xy=(1, TSTAR4 / 3600), xytext=(0, TSTAR3 / 3600),
            arrowprops=dict(arrowstyle='<->', color='#334155', lw=1.6))
ax.text(0.5, 55.2, '缩短 %.2f h\n（约 %.0f%%）'
        % (TSTAR3 / 3600 - TSTAR4 / 3600,
           100 * (TSTAR3 - TSTAR4) / TSTAR3),
        ha='center', fontsize=10, color='#334155')
ax.set_ylabel('烘干时长 / h'); ax.set_ylim(0, 65)
ax.set_title('(b) 收缩对烘干时长的影响', fontsize=12, fontweight='bold')

fig.tight_layout(rect=[0, 0, 1, 0.95])
p3 = os.path.join(HERE, 'fig3_validation.png')
fig.savefig(p3); plt.close(fig)
print('已保存', p3)

print('\n全部完成，共 3 张图 ->', HERE)
